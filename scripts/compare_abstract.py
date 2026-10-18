import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc_orig = docx.Document(r'C:\Users\risha\Downloads\Project Abstract Guideline_ORIGINAL_BACKUP.docx')
doc_done = docx.Document(r'C:\Users\risha\Downloads\AutoVision_RISC-V_Project_Abstract.docx')

print('=== ORIGINAL ABSTRACT TEMPLATE PARAGRAPHS ===')
for i, p in enumerate(doc_orig.paragraphs):
    if p.text.strip():
        print(f'P{i:02d}: [{p.style.name}] {p.text}')

print('\n=== COMPLETED ABSTRACT PARAGRAPHS ===')
for i, p in enumerate(doc_done.paragraphs):
    if p.text.strip():
        print(f'P{i:02d}: [{p.style.name}] {p.text}')
