import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
try:
    doc_t = docx.Document(r'C:\Users\risha\Downloads\Week 9 Lab Assignment Debugging, Tracing, and Profiling (1)_blank_template.docx')
    doc_s = docx.Document(r'C:\Users\risha\Downloads\Week 9 Lab Assignment Debugging, Tracing, and Profiling - Rishabh Thakur (250101084) Submission.docx')
    print('Week 9 Template paragraphs:', len(doc_t.paragraphs))
    print('Week 9 Submission paragraphs:', len(doc_s.paragraphs))
    for i in range(min(10, len(doc_t.paragraphs))):
        print(f'T{i:02d}: {doc_t.paragraphs[i].text[:50]}')
        print(f'S{i:02d}: {doc_s.paragraphs[i].text[:50]}')
except Exception as e:
    print('Error:', e)
