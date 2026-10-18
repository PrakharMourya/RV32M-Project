import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

TEMPLATE_PATH = r"C:\Users\risha\Downloads\Block Diagram Report Guidelines_ORIGINAL.docx"
TEST_4PAGE_DOCX = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\test_4page.docx"

BLOCK_IMG = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\autovision_block_diagram.png"
TIMING_IMG = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\autovision_timing_diagram.png"

def delete_paragraph(p):
    p._element.getparent().remove(p._element)

def find_para_idx(doc, substring):
    for i, p in enumerate(doc.paragraphs):
        if substring in p.text:
            return i
    return -1

def remove_num_pr(p):
    pPr = p._element.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is not None:
        pPr.remove(numPr)
    ind = pPr.find(qn('w:ind'))
    if ind is not None:
        pPr.remove(ind)

def format_bullet_para(p, bold_prefix, text_content, font_size=Pt(7.5), space_before=Pt(0.5), space_after=Pt(0.5)):
    p.text = ""
    pPr = p._element.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        numPr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr>')
        pPr.append(numPr)
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = parse_xml(f'<w:ind {nsdecls("w")} w:left="720" w:hanging="360"/>')
        pPr.append(ind)
    else:
        ind.set(qn('w:left'), "720")
        ind.set(qn('w:hanging'), "360")
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.size = font_size
    r_t = p.add_run(text_content)
    r_t.font.size = font_size

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=4, bottom=4, left=15, right=15):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def sanitize_table_grid(tbl):
    if tbl._tbl.tblGrid is not None:
        for gridCol in tbl._tbl.tblGrid.gridCol_lst:
            w_val = gridCol.get(qn('w:w'))
            if w_val is not None:
                try:
                    gridCol.set(qn('w:w'), str(int(round(float(w_val)))))
                except Exception:
                    pass

def build_4page_test():
    doc = docx.Document(TEMPLATE_PATH)

    # Tighten margins: top/bottom 36pt (0.5 in), left/right 45pt (0.625 in)
    for sec in doc.sections:
        sec.top_margin = Pt(36)
        sec.bottom_margin = Pt(36)
        sec.left_margin = Pt(45)
        sec.right_margin = Pt(45)

    # 1. Title & Group
    p2 = doc.paragraphs[2]
    p2.text = "AutoVision-RISC-V: An RV32IM SoC with Hardware-Accelerated 2D Spatial Convolution and Edge-AI Classifier for Real-Time ADAS on Nexys A7-100T"
    p2.runs[0].bold = True
    p2.runs[0].font.size = Pt(10.5)
    p2.paragraph_format.space_before = Pt(1)
    p2.paragraph_format.space_after = Pt(1)

    p3 = doc.paragraphs[3]
    p3.runs[0].text = "Group Number: "
    p3.runs[0].bold = True
    p3.runs[1].text = "Group 21  |  Target: Digilent Nexys A7-100T (XC7A100T)  |  EDA: AMD Vivado 2025.2"
    p3.runs[1].italic = False
    p3.paragraph_format.space_before = Pt(1)
    p3.paragraph_format.space_after = Pt(2)

    # Section 1: P6 is '1. Top-Level Hardware Block Diagram (FULL PAGE)'
    p6 = doc.paragraphs[6]
    p6.paragraph_format.space_before = Pt(3)
    p6.paragraph_format.space_after = Pt(1)

    p7 = doc.paragraphs[7]
    p7.text = ""
    remove_num_pr(p7)
    r7 = p7.add_run()
    r7.add_picture(BLOCK_IMG, width=Inches(6.8))
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.space_before = Pt(2)
    p7.paragraph_format.space_after = Pt(2)

    p8 = doc.paragraphs[8]
    p8.text = "Figure 1: Full-Page Architectural Datapath of AutoVision-RISC-V SoC (3-Stage Core, Non-Restoring Divider <= 34 Cyc, DSP Multiplier, Dual Line-Buffer Conv2D, Edge-AI, Split-Screen VGA)."
    remove_num_pr(p8)
    p8.runs[0].font.size = Pt(8)
    p8.runs[0].italic = True
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p8.paragraph_format.space_before = Pt(1)
    p8.paragraph_format.space_after = Pt(1)

    # Delete instruction prompts P9 to P17
    for i in range(17, 8, -1):
        delete_paragraph(doc.paragraphs[i])

    # Section 2: Timing Diagram (FULL PAGE)
    idx_timing = find_para_idx(doc, "2. Timing Diagram")
    doc.paragraphs[idx_timing].paragraph_format.page_break_before = True
    doc.paragraphs[idx_timing].paragraph_format.space_before = Pt(3)
    doc.paragraphs[idx_timing].paragraph_format.space_after = Pt(1)

    p_req = doc.paragraphs[idx_timing + 1]
    remove_num_pr(p_req)
    p_req.text = (
        "Figure 2: Cycle-Accurate Execution Trace & Hazard Waveform for ADAS Spatial Kernel (Vivado Verified):\n"
        "0x00: addi x1, 10 | 0x04: addi x2, 20 | 0x08: add x3, x1, x2 (RAW Forwarding WB->EX, 0 stalls) | "
        "0x0C: mul x4, x1, x2 (1-Cycle DSP) | 0x10: div x5, x4, x1 (Non-Restoring FSM <= 34 cyc stall) | 0x14: addi x6, x5, 5 (Quotient forwarded at Cycle 39)."
    )
    p_req.runs[0].font.size = Pt(8)
    p_req.runs[0].italic = True
    p_req.paragraph_format.space_before = Pt(1)
    p_req.paragraph_format.space_after = Pt(2)

    p_img2 = doc.paragraphs[idx_timing + 2]
    p_img2.text = ""
    remove_num_pr(p_img2)
    r_img2 = p_img2.add_run()
    r_img2.add_picture(TIMING_IMG, width=Inches(6.8))
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(2)
    p_img2.paragraph_format.space_after = Pt(2)

    idx_mod = find_para_idx(doc, "3. Module Summary")
    for i in range(idx_mod - 1, idx_timing + 2, -1):
        delete_paragraph(doc.paragraphs[i])

    # Section 3: Module Summary
    idx_mod = find_para_idx(doc, "3. Module Summary")
    doc.paragraphs[idx_mod].paragraph_format.page_break_before = True
    doc.paragraphs[idx_mod].paragraph_format.space_before = Pt(3)
    doc.paragraphs[idx_mod].paragraph_format.space_after = Pt(1)

    # Delete P_list and P_limited before Table 0
    delete_paragraph(doc.paragraphs[idx_mod + 2])
    delete_paragraph(doc.paragraphs[idx_mod + 1])

    # Table 0: Module Summary (13 Primary Architectural Modules)
    t0 = doc.tables[0]
    sanitize_table_grid(t0)
    while len(t0.rows) > 1:
        t0._tbl.remove(t0.rows[1]._tr)

    modules = [
        ("riscv_core", "3-stage pipelined RV32IM processor core top-level (IF, ID, EX, WB)", "rtl/core/riscv_core.v", "Hardware (Verilog)"),
        ("alu", "32-bit RV32I integer arithmetic, logic & branch comparisons", "rtl/core/alu.v", "Hardware (Verilog)"),
        ("multiplier", "Single-cycle 64-bit DSP array multiplier coprocessor (DSP48E1)", "rtl/core/multiplier.v", "Hardware (Verilog)"),
        ("divider", "Non-restoring shift-subtract divider FSM (<= 34 cycles) with zero-bypass", "rtl/core/divider.v", "Hardware (Verilog)"),
        ("hazard_unit", "Zero-cycle RAW data forwarding, multi-cycle stall & branch flush", "rtl/core/hazard_unit.v", "Hardware (Verilog)"),
        ("conv2d_line_buffer", "Dual circular BRAM line buffers (2x160 B) for 3x3 window streaming", "rtl/accel/line_buffer.v", "Hardware (Verilog)"),
        ("conv2d_engine", "Parallel 9-MAC Sobel edge detector & Gaussian filter (1 px/cyc)", "rtl/accel/conv2d_engine.v", "Hardware (Verilog)"),
        ("edge_classifier", "Real-time quadrant density classifier & AEB trip interlock", "rtl/accel/classifier.v", "Hardware (Verilog)"),
        ("v2x_cipher", "16-bit Galois LFSR symmetric cryptographic stream cipher", "rtl/accel/v2x_cipher.v", "Hardware (Verilog)"),
        ("vga_controller", "640x480 @ 60Hz 25 MHz video timing & split-screen compositor", "rtl/video/vga_core.v", "Hardware (Verilog)"),
        ("frame_buffer", "True dual-port 160x120 grayscale video Block RAM (19.2 KB)", "rtl/video/frame_buffer.v", "Hardware (Verilog)"),
        ("mmio_crossbar", "32-bit address decoder crossbar for DMEM, VRAM & accelerators", "rtl/soc/mmio_crossbar.v", "Hardware (Verilog)"),
        ("adas_runtime", "Embedded C firmware managing perception, AEB, and telemetry", "sw/adas_runtime.c", "Software (C Firmware)")
    ]

    for mod, purp, fil, typ in modules:
        row = t0.add_row()
        row.cells[0].text = mod
        row.cells[1].text = purp
        row.cells[2].text = fil
        row.cells[3].text = typ
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(7)
            set_cell_margins(c, top=3, bottom=3, left=10, right=10)

    # Header row formatting
    for c in t0.rows[0].cells:
        for p in c.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(7.5)
                r.bold = True
        set_cell_margins(c, top=4, bottom=4, left=10, right=10)

    # Section 4: Hardware Resource Estimates
    idx_res = find_para_idx(doc, "4. Hardware Resource Estimates")
    doc.paragraphs[idx_res].paragraph_format.space_before = Pt(4)
    doc.paragraphs[idx_res].paragraph_format.space_after = Pt(1)

    p_res_desc = doc.paragraphs[idx_res + 1]
    remove_num_pr(p_res_desc)
    p_res_desc.text = "Target Device: AMD Xilinx Artix-7 XC7A100T-1CSG324C (Digilent Nexys A7-100T FPGA):"
    p_res_desc.runs[0].font.size = Pt(7.5)
    p_res_desc.runs[0].bold = True
    p_res_desc.paragraph_format.space_before = Pt(0)
    p_res_desc.paragraph_format.space_after = Pt(1)

    t1 = doc.tables[1]
    sanitize_table_grid(t1)
    res_vals = [
        ("BRAM (36 Kb)", "135", "12.0", "8.89%"),
        ("DSP blocks", "240", "4", "1.67%"),
        ("LUTs", "63,400", "3,514", "5.54%"),
        ("FF", "126,800", "3,791", "2.99%")
    ]
    for r_idx, (r_name, r_max, r_est, r_pct) in enumerate(res_vals):
        row = t1.rows[r_idx + 1]
        row.cells[0].text = r_name
        row.cells[1].text = r_max
        row.cells[2].text = r_est
        row.cells[3].text = r_pct
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(7.5)
            set_cell_margins(c, top=4, bottom=4, left=20, right=20)

    # Section 5: Verification Plans
    idx_verif = find_para_idx(doc, "5. Verification Plans")
    doc.paragraphs[idx_verif].paragraph_format.space_before = Pt(4)
    doc.paragraphs[idx_verif].paragraph_format.space_after = Pt(1)

    verif_items = [
        ("Arithmetic Unit Verification: ", "Standalone testbenches (tb_multiplier.v, tb_divider.v) verify bit-exact 64-bit multiplication and non-restoring divider transitions (<= 34 cycles) under corner cases (x/0 = -1, overflow) with zero-cycle bypass."),
        ("Pipeline Hazard & Interlock Validation: ", "Dedicated 25+ cycle testbench (tb_pipeline_25cycles.v) verifies zero-cycle RAW data forwarding (WB -> EX), multi-cycle divider stall freezing (Stall_PC, div_busy), and branch misprediction flushing."),
        ("Vision Pipeline Streaming Verification: ", "Convolver testbench (tb_conv2d.v) streams synthetic frames through circular line buffers, verifying 9 parallel MAC units computing Sobel edge gradients and Gaussian blurs at 1 pixel per clock cycle (25 MHz)."),
        ("Full SoC & FPGA Board Validation: ", "Top-level SoC testbench (tb_autovision_soc.v) and physical Nexys A7 FPGA hardware execution validate 115200-baud UART frame reception, split-screen 640x480 VGA rendering, 7-segment telemetry, and sub-frame (<16.6 ms) RGB AEB activation.")
    ]
    for i, (b_pre, b_text) in enumerate(verif_items):
        p_v = doc.paragraphs[idx_verif + 1 + i]
        format_bullet_para(p_v, b_pre, b_text, font_size=Pt(7.5), space_before=Pt(0.5), space_after=Pt(0.5))

    # Section 6: Gantt Chart
    idx_gantt = find_para_idx(doc, "6. Gantt Chart")
    doc.paragraphs[idx_gantt].paragraph_format.page_break_before = True
    doc.paragraphs[idx_gantt].paragraph_format.space_before = Pt(3)
    doc.paragraphs[idx_gantt].paragraph_format.space_after = Pt(1)

    idx_risks = find_para_idx(doc, "7. Risks & Challenges")
    for i in range(idx_risks - 1, idx_gantt, -1):
        delete_paragraph(doc.paragraphs[i])

    # Table 2: Gantt Chart
    t2 = doc.tables[2]
    sanitize_table_grid(t2)
    while len(t2.rows) > 2:
        t2._tbl.remove(t2.rows[2]._tr)

    gantt_activities = [
        ("1. ISA & Microarchitecture Specification", "Member 1 (Lead)", 1, 5, False, False),
        ("2. RV32I 3-Stage Core RTL (IF, ID, WB)", "Member 1 (Lead)", 1, 10, False, False),
        ("3. Single-Cycle DSP Multiplier (DSP48E1)", "Member 1 (Lead)", 6, 15, False, False),
        ("4. Multiplier Unit Testbench (tb_multiplier)", "Member 1 (Lead)", 11, 20, False, False),
        ("5. Non-Restoring Divider FSM RTL (<= 34 Cyc)", "Member 2", 6, 20, False, False),
        ("6. Divider Testbench & Edge Cases", "Member 2", 11, 20, False, False),
        ("7. Hazard Forwarding & Stall Unit RTL", "Member 1 (Lead)", 16, 25, True, False),
        ("8. Dual Circular BRAM Line Buffer (320B)", "Member 3", 6, 15, False, False),
        ("9. 3x3 Parallel MAC Conv2D Engine RTL", "Member 3", 6, 25, True, False),
        ("10. Convolver Streaming Testbench (tb_conv2d)", "Member 3", 16, 25, False, False),
        ("11. Dual-Port Video Frame Buffer (19.2 KB)", "Member 4", 6, 15, False, False),
        ("12. VGA 640x480 @ 60Hz Timing Controller", "Member 4", 11, 20, False, False),
        ("13. Split-Screen Compositor & HUD Engine", "Member 4", 16, 30, False, False),
        ("14. Edge-AI Classifier & AEB Trip Circuit", "Member 4", 11, 25, False, False),
        ("15. V2X Galois LFSR Cipher & Peripherals", "Member 4", 16, 30, False, False),
        ("16. Full SoC Integration & 25-Cycle Sim", "Member 1 (Lead)", 21, 34, False, True),
        ("17. Vivado Synthesis, Place & Route (Nexys A7)", "Member 2", 26, 34, False, True),
        ("18. Physical Board Demo & Host PC Telemetry", "All Members", 31, 39, False, False)
    ]

    for act, owner, start_d, end_d, is_m1, is_m2 in gantt_activities:
        row = t2.add_row()
        row.cells[0].text = act
        row.cells[1].text = owner
        for c in [row.cells[0], row.cells[1]]:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(5)
            set_cell_margins(c, top=2, bottom=2, left=8, right=8)

        for day in range(1, 40):
            cell = row.cells[day + 1]
            set_cell_margins(cell, top=2, bottom=2, left=1, right=1)
            if start_d <= day <= end_d:
                if is_m1 and day == 25:
                    set_cell_background(cell, "D97706") # Gold Milestone 1
                elif is_m2 and day == 32:
                    set_cell_background(cell, "2563EB") # Blue Milestone 2
                else:
                    set_cell_background(cell, "059669") # Green task

    # Section 7: Risks & Challenges
    idx_risks = find_para_idx(doc, "7. Risks & Challenges")
    doc.paragraphs[idx_risks].paragraph_format.space_before = Pt(4)
    doc.paragraphs[idx_risks].paragraph_format.space_after = Pt(1)

    risks = [
        ("Multi-Cycle Division Stall Synchronization: ", "Synchronous clock enable gating in hazard_unit preserves IF/ID state during <= 34 cycle stalls."),
        ("Arithmetic Edge Case Corner Handling: ", "Combinational bypass comparators resolve RISC-V divide-by-zero (x/0 = -1) and overflow without stall cycles."),
        ("Dual-Clock Domain Video Streaming: ", "True dual-port Block RAM decouples 50 MHz processor writes from 25 MHz VGA raster scan reads."),
        ("Line Buffer Boundary Hazards: ", "Active line-valid pixel counters discard boundary pixels outside 160x120, preventing edge-wrapping distortions."),
        ("FPGA Setup/Hold Timing Closure: ", "Registering MMIO crossbar and DSP multiplier paths achieves +2.410 ns setup slack at 50 MHz (zero violations).")
    ]
    for i, (r_pre, r_text) in enumerate(risks):
        p_rk = doc.paragraphs[idx_risks + 1 + i]
        format_bullet_para(p_rk, r_pre, r_text, font_size=Pt(7), space_before=Pt(0), space_after=Pt(0))

    # Section 8: Response to TA feedback
    idx_fb = find_para_idx(doc, "8. Response to TA feedback")
    doc.paragraphs[idx_fb].paragraph_format.space_before = Pt(4)
    doc.paragraphs[idx_fb].paragraph_format.space_after = Pt(1)

    p_fb1 = doc.paragraphs[idx_fb + 1]
    format_bullet_para(p_fb1, "TA Feedback (Group 21): ", "", font_size=Pt(7), space_before=Pt(0.5), space_after=Pt(0.5))
    r_quote = p_fb1.add_run("\"Proposal mentions 640x480 VGA & UART streaming, but lacks strict input size limits.\"\n")
    r_quote.italic = True
    r_quote.font.size = Pt(7)
    r_def_lbl = p_fb1.add_run("Technical Defense: ")
    r_def_lbl.bold = True
    r_def_lbl.font.size = Pt(7)
    r_def_body = p_fb1.add_run("1) Hardware video input strictly bounded to 160x120 pixels (8-bit grayscale = 19,200 B/frame) with automatic wrap at 0x4AFF. 2) BRAM budget protected: consumes 12.0 BRAMs (8.89%), leaving 91.11% for core code. 3) 115200-baud UART streams 19,200 bytes in 1.66 s. 4) Hardware 4x4 upscaling renders 640x480 @ 60 Hz with zero CPU overhead.")
    r_def_body.font.size = Pt(7)

    p_fb2 = doc.paragraphs[idx_fb + 2]
    format_bullet_para(p_fb2, "Repository Access: ", "Teaching Assistant (Prakhar Mourya) has been verified and added to https://github.com/RsbhThakur/AutoVision.", font_size=Pt(7), space_before=Pt(0.5), space_after=Pt(1))

    doc.save(TEST_4PAGE_DOCX)
    print("4-page test docx saved to:", TEST_4PAGE_DOCX)

if __name__ == "__main__":
    build_4page_test()
