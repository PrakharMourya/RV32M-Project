import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import shutil

TEMPLATE_PATH = r"C:\Users\risha\Downloads\Block Diagram Report Guidelines_ORIGINAL.docx"
TEST_DOCX_PATH = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\test_clean.docx"
BLOCK_IMG = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\autovision_block_diagram.png"
TIMING_IMG = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\autovision_timing_diagram.png"

def test_build():
    doc = docx.Document(TEMPLATE_PATH)
    
    # Check title replacement
    # P2: "Concise, specific, and technical"
    p2 = doc.paragraphs[2]
    p2.text = "AutoVision-RISC-V: An RV32IM SoC with Hardware-Accelerated 2D Spatial Convolution and Edge-AI Classifier for Real-Time ADAS on Nexys A7-100T"
    p2.runs[0].font.bold = True
    p2.runs[0].font.italic = False
    
    # P3: "Group Number: Add here"
    p3 = doc.paragraphs[3]
    p3.text = "Group Number: Group 21"
    p3.runs[0].font.bold = True
    
    # P16 is empty paragraph under Section 1
    p16 = doc.paragraphs[16]
    run1 = p16.add_run()
    run1.add_picture(BLOCK_IMG, width=Inches(6.3))
    p16.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # P27 is empty paragraph under Section 2
    p27 = doc.paragraphs[27]
    run2 = p27.add_run()
    run2.add_picture(TIMING_IMG, width=Inches(6.3))
    p27.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.save(TEST_DOCX_PATH)
    print("Saved test docx:", TEST_DOCX_PATH)

if __name__ == "__main__":
    test_build()
