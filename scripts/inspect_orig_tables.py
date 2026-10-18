import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document(r'C:\Users\risha\Downloads\Block Diagram Report Guidelines_ORIGINAL.docx')

for t_idx, t in enumerate(doc.tables):
    print(f'\n=== Table {t_idx} ===')
    print('tblStyle:', t.style.name if t.style else None)
    print('tblPr XML:', t._tbl.tblPr.xml[:300] if t._tbl.tblPr is not None else None)
    for r_idx, row in enumerate(t.rows[:2]):
        print(f'Row {r_idx}:')
        for c_idx, cell in enumerate(row.cells[:4]):
            p = cell.paragraphs[0]
            r_info = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
            print(f'  Cell ({r_idx},{c_idx}): text="{cell.text.strip()}", p_style="{p.style.name}", runs={r_info}')
