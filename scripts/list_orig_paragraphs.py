import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document(r'C:\Users\risha\Downloads\Block Diagram Report Guidelines_ORIGINAL.docx')

print(f'Total paragraphs: {len(doc.paragraphs)}')
for i, p in enumerate(doc.paragraphs):
    print(f'P{i:02d} [{p.style.name}]: "{p.text}"')

print(f'\nTotal tables: {len(doc.tables)}')
for t_idx, t in enumerate(doc.tables):
    print(f'Table {t_idx}: {len(t.rows)} rows, {len(t.columns)} cols')
