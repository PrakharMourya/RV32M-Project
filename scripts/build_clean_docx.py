import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SRC_TEMPLATE = r"C:\Users\risha\Downloads\Block Diagram Report Guidelines_ORIGINAL.docx"
OUT_DOCX = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\Group21_Block_Diagram_Report.docx"
DOWNLOAD_DOCX = r"C:\Users\risha\Downloads\Group21_Block_Diagram_Report.docx"
TARGET_GUIDELINE = r"C:\Users\risha\Downloads\Block Diagram Report Guidelines.docx"

BLOCK_IMG = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\autovision_block_diagram.png"
TIMING_IMG = r"C:\Users\risha\OneDrive\Desktop\AutoVision\docs\autovision_timing_diagram.png"

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def sanitize_table_grid(tbl):
    if tbl._tbl.tblGrid is not None:
        for gridCol in tbl._tbl.tblGrid.gridCol_lst:
            w_val = gridCol.get(qn('w:w'))
            if w_val is not None:
                try:
                    gridCol.set(qn('w:w'), str(int(round(float(w_val)))))
                except Exception:
                    pass

def build_report():
    doc = docx.Document(SRC_TEMPLATE)

    # 1. Title
    # P1: "Title"
    # P2: "Concise, specific, and technical" -> replace with our Title
    p2 = doc.paragraphs[2]
    p2.text = "AutoVision-RISC-V: An RV32IM SoC with Hardware-Accelerated 2D Spatial Convolution and Edge-AI Classifier for Real-Time ADAS on Nexys A7-100T"
    for r in p2.runs:
        r.bold = True
        r.italic = False

    # P3: "Group Number: Add here"
    # Run 0 is "Group Number: ", Run 1 is "Add here"
    p3 = doc.paragraphs[3]
    p3.runs[0].text = "Group Number: "
    p3.runs[0].bold = True
    p3.runs[1].text = "Group 21  |  Target FPGA: Digilent Nexys A7-100T (XC7A100T)  |  EDA: AMD Vivado 2025.2"
    p3.runs[1].italic = False
    p3.runs[1].bold = False

    # 2. Section 1: Top-Level Hardware Block Diagram (FULL PAGE)
    # P16 is empty paragraph under Section 1
    p16 = doc.paragraphs[16]
    r16 = p16.add_run()
    r16.add_picture(BLOCK_IMG, width=Inches(6.25))
    p16.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Caption in P17
    p17 = doc.paragraphs[17]
    p17.text = "Figure 1: Full-Page Top-Level Architectural Datapath of AutoVision-RISC-V SoC showing 3-stage RV32IM core, non-restoring divider FSM (<= 34 cycles), single-cycle DSP multiplier, dual circular line buffers, 2D Conv2D engine, Edge-AI classifier, and split-screen VGA compositor."
    p17.runs[0].italic = True
    p17.runs[0].font.size = Pt(8.5)

    # 3. Section 2: Timing Diagram (FULL PAGE)
    # P27 is empty paragraph under Section 2
    p27 = doc.paragraphs[27]
    r27 = p27.add_run()
    r27.add_picture(TIMING_IMG, width=Inches(6.25))
    p27.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add assembly trace note in P26
    p26 = doc.paragraphs[26]
    p26.text = (
        "Figure 2: Cycle-Accurate Execution Trace & Hazard Waveform for ADAS Spatial Kernel:\n"
        "0x00: addi x1, 10 | 0x04: addi x2, 20 | 0x08: add x3, x1, x2 (RAW Forwarding WB->EX, 0 stalls) | "
        "0x0C: mul x4, x1, x2 (1-Cycle DSP) | 0x10: div x5, x4, x1 (Non-Restoring FSM <= 34 cycle stall) | 0x14: addi x6, x5, 5 (Quotient forwarded at Cycle 39)."
    )
    p26.runs[0].italic = True
    p26.runs[0].font.size = Pt(8.5)

    # 4. Section 3: Module Summary (Table Only)
    t0 = doc.tables[0]
    sanitize_table_grid(t0)
    # Clear existing rows 1 and 2
    while len(t0.rows) > 1:
        t0._tbl.remove(t0.rows[1]._tr)

    modules = [
        ("riscv_core", "3-stage pipelined RV32IM processor core top-level", "rtl/core/riscv_core.v", "Hardware (Verilog)"),
        ("if_stage", "Instruction Fetch, PC generator, branch target multiplexer", "rtl/core/if_stage.v", "Hardware (Verilog)"),
        ("id_ex_stage", "Instruction decode, 32x32 register file & sign extension", "rtl/core/id_ex_stage.v", "Hardware (Verilog)"),
        ("alu", "32-bit RV32I integer arithmetic, logic & branch comparisons", "rtl/core/alu.v", "Hardware (Verilog)"),
        ("multiplier", "Single-cycle 64-bit DSP array multiplier coprocessor (DSP48E1)", "rtl/core/multiplier.v", "Hardware (Verilog)"),
        ("divider", "Non-restoring shift-subtract divider FSM (<= 34 cycles) with zero-bypass", "rtl/core/divider.v", "Hardware (Verilog)"),
        ("hazard_unit", "Zero-cycle RAW forwarding, multi-cycle stall & branch flush", "rtl/core/hazard_unit.v", "Hardware (Verilog)"),
        ("wb_stage", "Writeback multiplexer & register commit interface", "rtl/core/wb_stage.v", "Hardware (Verilog)"),
        ("conv2d_line_buffer", "Dual circular BRAM line buffers (2x160 B) for 3x3 windowing", "rtl/accel/line_buffer.v", "Hardware (Verilog)"),
        ("conv2d_engine", "Parallel 9-MAC Sobel edge detector & Gaussian filter (1 px/cyc)", "rtl/accel/conv2d_engine.v", "Hardware (Verilog)"),
        ("edge_classifier", "Real-time quadrant density classifier & AEB trip interlock", "rtl/accel/classifier.v", "Hardware (Verilog)"),
        ("v2x_cipher", "16-bit Galois LFSR symmetric cryptographic stream cipher", "rtl/accel/v2x_cipher.v", "Hardware (Verilog)"),
        ("vga_controller", "640x480 @ 60Hz 25 MHz video timing generator & DAC interface", "rtl/video/vga_core.v", "Hardware (Verilog)"),
        ("frame_buffer", "True dual-port 160x120 grayscale video Block RAM (19.2 KB)", "rtl/video/frame_buffer.v", "Hardware (Verilog)"),
        ("split_screen_mux", "Real-time split-screen compositor: raw feed, laser line, Sobel, HUD", "rtl/video/split_mux.v", "Hardware (Verilog)"),
        ("mmio_crossbar", "32-bit address decoder crossbar for DMEM, VRAM & accelerators", "rtl/soc/mmio_crossbar.v", "Hardware (Verilog)"),
        ("uart_transceiver", "Full-duplex 115200-baud UART transceiver with RX/TX FIFOs", "rtl/soc/uart_transceiver.v", "Hardware (Verilog)"),
        ("seven_seg_driver", "8-digit multiplexed 7-segment display driver (1 kHz refresh)", "rtl/soc/seven_seg.v", "Hardware (Verilog)"),
        ("actuator_controller", "Tri-color RGB brake actuators (AEB) & GPIO switches/LEDs", "rtl/soc/actuators.v", "Hardware (Verilog)"),
        ("autovision_soc", "Top-level SoC interconnecting core, BRAM, video & peripherals", "rtl/soc/autovision_soc.v", "Hardware (Verilog)"),
        ("adas_runtime", "Embedded C firmware managing perception, AEB, and telemetry", "sw/adas_runtime.c", "Software (C Firmware)")
    ]

    for mod, purp, fil, typ in modules:
        row = t0.add_row()
        row.cells[0].text = mod
        row.cells[1].text = purp
        row.cells[2].text = fil
        row.cells[3].text = typ
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
            set_cell_margins(c, top=20, bottom=20, left=40, right=40)

    # 5. Section 4: Hardware Resource Estimates
    t1 = doc.tables[1]
    sanitize_table_grid(t1)
    res_vals = [
        ("BRAM", "135", "12.0", "8.89%"),
        ("DSP blocks", "240", "4", "1.67%"),
        ("LUTs", "63,400", "3,514", "5.54%"),
        ("FF", "126,800", "3,791", "2.99%")
    ]
    for r_idx, (r_name, r_max, r_est, r_pct) in enumerate(res_vals):
        row = t1.rows[r_idx + 1]
        row.cells[0].text = r_name
        row.cells[1].text = r_max
        row.cells[2].text = r_est
        row.cells[3].text = r_pct
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
            set_cell_margins(c, top=25, bottom=25, left=50, right=50)

    # 6. Section 5: Verification Plans (≤ 6–8 lines total OR 3–4 bullets)
    # P36 is 'Examples:' -> replace with actual verification bullets
    doc.paragraphs[36].text = "• Arithmetic Unit Verification: Standalone testbenches (tb_multiplier.v, tb_divider.v) verify bit-exact 64-bit multiplication (MUL/H/HSU/HU) and validate non-restoring divider state transitions (<= 34 cycles) under corner cases (x/0 = -1, overflow) with zero-cycle bypass."
    doc.paragraphs[36].runs[0].font.size = Pt(8.5)

    doc.paragraphs[37].text = "• Pipeline Hazard & Interlock Validation: Dedicated 25+ cycle testbench (tb_pipeline_25cycles.v) verifies zero-cycle RAW data forwarding (WB -> EX), multi-cycle non-restoring divider stall freezing (Stall_PC, div_busy), and branch misprediction flushing."
    doc.paragraphs[37].runs[0].font.size = Pt(8.5)

    doc.paragraphs[38].text = "• Vision Pipeline Streaming Verification: Convolver testbench (tb_conv2d.v) streams synthetic frames through circular line buffers, verifying that 9 parallel MAC units compute Sobel edge gradients and Gaussian blurs at 1 pixel per clock cycle (25 MHz)."
    doc.paragraphs[38].runs[0].font.size = Pt(8.5)

    doc.paragraphs[39].text = "• Full SoC & FPGA Board Validation: Top-level SoC testbench (tb_autovision_soc.v) and physical Nexys A7 FPGA hardware execution validate 115200-baud UART frame reception, split-screen 640x480 VGA rendering, 7-segment display telemetry, and sub-frame (<16.6 ms) RGB AEB activation."
    doc.paragraphs[39].runs[0].font.size = Pt(8.5)

    # 7. Section 6: Gantt Chart
    t2 = doc.tables[2]
    sanitize_table_grid(t2)
    # Keep header row 0 and row 1, remove rows 2 to 6
    while len(t2.rows) > 2:
        t2._tbl.remove(t2.rows[2]._tr)

    gantt_activities = [
        ("1. ISA & Microarchitecture Specification", "Member 1 (Lead)", 1, 5, False, False),
        ("2. RV32I 3-Stage Core RTL (IF, ID, WB)", "Member 1 (Lead)", 1, 10, False, False),
        ("3. Single-Cycle DSP Multiplier (DSP48E1)", "Member 1 (Lead)", 6, 15, False, False),
        ("4. Multiplier Unit Testbench (tb_multiplier)", "Member 1 (Lead)", 11, 20, False, False),
        ("5. Non-Restoring Divider FSM RTL (<= 34 Cyc)", "Member 2", 6, 20, False, False),
        ("6. Divider Testbench & Edge Cases", "Member 2", 11, 20, False, False),
        ("7. Hazard Forwarding & Stall Unit RTL", "Member 1 (Lead)", 16, 25, True, False),   # M1
        ("8. Dual Circular BRAM Line Buffer (320B)", "Member 3", 6, 15, False, False),
        ("9. 3x3 Parallel MAC Conv2D Engine RTL", "Member 3", 11, 25, True, False),   # M1
        ("10. Convolver Streaming Testbench (tb_conv2d)", "Member 3", 16, 25, False, False),
        ("11. Dual-Port Video Frame Buffer (19.2 KB)", "Member 4", 6, 15, False, False),
        ("12. VGA 640x480 @ 60Hz Timing Controller", "Member 4", 11, 20, False, False),
        ("13. Split-Screen Compositor & HUD Engine", "Member 4", 16, 31, False, False),
        ("14. Edge-AI Classifier & AEB Trip Circuit", "Member 4", 11, 25, False, False),
        ("15. V2X Galois LFSR Cipher & Peripherals", "Member 4", 16, 31, False, False),
        ("16. Full SoC Integration & 25-Cycle Sim", "Member 1 (Lead)", 21, 35, False, True),  # M2
        ("17. Vivado Synthesis, Place & Route (Nexys A7)", "Member 2", 26, 35, False, True), # M2
        ("18. Physical Board Demo & Host PC Telemetry", "All Members", 32, 39, False, False) # Nov 1-8
    ]

    for act_name, owner, start_d, end_d, is_m1, is_m2 in gantt_activities:
        row = t2.add_row()
        row.cells[0].text = act_name
        row.cells[1].text = owner
        for c in [row.cells[0], row.cells[1]]:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(6.5)
            set_cell_margins(c, top=15, bottom=15, left=25, right=25)

        for day in range(1, 40):
            cell = row.cells[day + 1]
            set_cell_margins(cell, top=15, bottom=15, left=8, right=8)
            if start_d <= day <= end_d:
                if is_m1 and day == 25:
                    set_cell_background(cell, "D97706") # Gold Milestone 1
                elif is_m2 and day == 32:
                    set_cell_background(cell, "2563EB") # Blue Milestone 2
                else:
                    set_cell_background(cell, "059669") # Green task

    # 8. Section 7: Risks & Challenges (≤ 5 bullets, 1 line each)
    doc.paragraphs[59].text = "1. Multi-Cycle Division Stall Synchronization: Synchronous clock enable gating in hazard_unit preserves IF/ID state and prevents instruction loss during <= 34 cycle stalls."
    doc.paragraphs[59].runs[0].font.size = Pt(8.5)

    doc.paragraphs[60].text = "2. Arithmetic Edge Case Corner Handling: Combinational bypass comparators resolve RISC-V divide-by-zero (x/0 = -1) and 2's complement overflow without stall cycles."
    doc.paragraphs[60].runs[0].font.size = Pt(8.5)

    doc.paragraphs[61].text = "3. Dual-Clock Domain Video Streaming: True dual-port Block RAM decouples 50 MHz processor writes from 25 MHz VGA raster scan reads, preventing race conditions."
    doc.paragraphs[61].runs[0].font.size = Pt(8.5)

    doc.paragraphs[62].text = "4. Line Buffer Boundary Hazards: Active line-valid pixel counters discard boundary pixels outside 160x120, preventing edge-wrapping distortions during 3x3 convolution."
    doc.paragraphs[62].runs[0].font.size = Pt(8.5)

    doc.paragraphs[63].text = "5. FPGA Setup/Hold Timing Closure: Registering MMIO crossbar and DSP multiplier paths achieves +2.410 ns setup slack at 50 MHz, ensuring zero timing violations."
    doc.paragraphs[63].runs[0].font.size = Pt(8.5)

    # 9. Section 8: Response to TA feedback (for the project abstract)
    doc.paragraphs[66].text = (
        "TA Feedback (Group 21): \"The proposal mentions a 640 × 480 VGA display resolution and host PC streaming over UART, but it does not specify strict input size limits for incoming image.\"\n"
        "Point-by-Point Technical Defense:\n"
        "1. Strict Ingestion Resolution (160 x 120 Grayscale): Hardware video input is strictly bounded to 160 x 120 pixels (8-bit grayscale = 19,200 bytes/frame) with automatic address counter wrapping at 0x4AFF.\n"
        "2. On-Chip BRAM Budget Protection: An uncompressed 640x480x8-bit frame requires 307,200 bytes (>50% of Artix-7 100T BRAM). The 160x120 budget consumes only 12.0 BRAMs (8.89%), preserving 91.11% of BRAM for core code and line buffers.\n"
        "3. UART Ingestion Throughput & Latency: At standard 115,200 baud (11,520 payload bytes/s), streaming 19,200 bytes requires 1.66 seconds per frame, matching realistic ADAS obstacle updates without buffer overruns.\n"
        "4. Hardware 4x4 Upscaling to 640 x 480 @ 60 Hz: The VGA controller upscales video in pure hardware via 4x4 pixel replication with zero CPU cycles and zero external DDR RAM dependencies."
    )
    doc.paragraphs[66].runs[0].font.size = Pt(8.5)

    doc.paragraphs[67].text = "Teaching Assistant (Prakhar Mourya) has been verified and added to the official Git repository at https://github.com/RsbhThakur/AutoVision."
    doc.paragraphs[67].runs[0].font.size = Pt(8.5)
    doc.paragraphs[67].runs[0].bold = True

    # Save to OUT_DOCX
    doc.save(OUT_DOCX)
    print("Report built successfully at:", OUT_DOCX)

    # Save to DOWNLOAD_DOCX
    doc.save(DOWNLOAD_DOCX)
    print("Saved copy to Downloads:", DOWNLOAD_DOCX)

    # Overwrite TARGET_GUIDELINE if not locked
    try:
        doc.save(TARGET_GUIDELINE)
        print("Enriched template at:", TARGET_GUIDELINE)
    except PermissionError:
        print("Note: TARGET_GUIDELINE is locked by Word.")

if __name__ == "__main__":
    build_report()
