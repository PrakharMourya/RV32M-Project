import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document(r'docs/Group21_Block_Diagram_Report.docx')

print('=== PARAGRAPHS IN MODIFIED DOCX ===')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() or len(p.runs) > 0:
        has_pic = 'PIC' if any('graphic' in r._r.xml or 'pict' in r._r.xml for r in p.runs) else ''
        print(f'P{i:02d}: style="{p.style.name}", text="{p.text[:60]}" {has_pic}')

print('\n=== TABLES IN MODIFIED DOCX ===')
for t_idx, tbl in enumerate(doc.tables):
    print(f'Table {t_idx}: rows={len(tbl.rows)}, cols={len(tbl.columns)}')
