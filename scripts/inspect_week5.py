import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
try:
    doc_t = docx.Document(r'C:\Users\risha\Downloads\Week_5_Laboratory_Assignment.docx')
    doc_s = docx.Document(r'C:\Users\risha\Downloads\Week_5_Laboratory_Assignment_Completed.docx')
    print('Week 5 Template paragraphs:', len(doc_t.paragraphs))
    print('Week 5 Submission paragraphs:', len(doc_s.paragraphs))
    print('Week 5 Template tables:', len(doc_t.tables))
    print('Week 5 Submission tables:', len(doc_s.tables))
    for i in range(min(15, len(doc_t.paragraphs))):
        print(f'T{i:02d}: {doc_t.paragraphs[i].text[:50]}')
        print(f'S{i:02d}: {doc_s.paragraphs[i].text[:50]}')
except Exception as e:
    print('Error:', e)
